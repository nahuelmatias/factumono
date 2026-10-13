"""Lectura determinística de facturas anteriores: PDF de ARCA («Comprobantes en línea») y los de
Factumono (mismo diseño), más el Excel/CSV de «Mis Comprobantes».

Nada de adivinar con IA: el PDF de ARCA tiene un diseño fijo, así que se lee el texto con sus
posiciones (pdfium) y se buscan las etiquetas («Razón Social:», «CAE N°:»...) y los valores que
tienen a la derecha o debajo. El QR (RG 4892) se decodifica y se usa para **validar** los datos
duros (CUIT, punto de venta, número, importe, CAE, documento del receptor): si no coinciden con
el texto, se avisa y no se registra nada como comprobante histórico.

Si el archivo no es de ARCA (o es una imagen sin texto) se devuelve `NoReconocido` con el texto
crudo, para que una IA lo interprete y lo cargue con `guardar_plantilla`.
"""

import base64
import contextlib
import csv
import io
import json
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from factumono.core import cuit as cuit_mod
from factumono.core.models import Concepto, DocTipo

QR_PREFIJO = "/fe/qr/?p="


# --- resultados -----------------------------------------------------------------------------


@dataclass
class ItemLeido:
    descripcion: str
    cantidad: Decimal
    unidad: str
    precio_unitario: Decimal
    codigo: str = "1"
    bonificacion_pct: Decimal = Decimal(0)
    bonificacion_importe: Decimal = Decimal(0)
    subtotal: Decimal = Decimal(0)


@dataclass
class FacturaLeida:
    origen: str  # 'arca' | 'factumono' | 'datos' (cargada por una IA)
    archivo: str = ""
    emisor_cuit: str = ""
    emisor_nombre: str = ""
    emisor_domicilio: str = ""
    emisor_condicion_iva: str = ""
    emisor_iibb: str = ""
    emisor_inicio_actividades: date | None = None
    cbte_tipo: int | None = None
    letra: str = ""
    punto_venta: int | None = None
    numero: int | None = None
    fecha: date | None = None
    receptor_doc_tipo: int = int(DocTipo.CONSUMIDOR_FINAL)
    receptor_doc_nro: int = 0
    receptor_nombre: str = ""
    receptor_condicion_iva: str = ""
    receptor_domicilio: str = ""
    concepto: Concepto = Concepto.PRODUCTOS
    periodo_desde: date | None = None
    periodo_hasta: date | None = None
    vto_pago: date | None = None
    condicion_venta: str = "Contado"
    items: list[ItemLeido] = field(default_factory=list)
    total: Decimal | None = None
    cae: str = ""
    cae_vto: date | None = None
    qr: dict[str, Any] | None = None
    advertencias: list[str] = field(default_factory=list)
    inconsistente: bool = False  # el QR no coincide con el texto

    @property
    def es_externo_completo(self) -> bool:
        """¿Tiene todo lo necesario para registrarse como comprobante histórico?"""
        return bool(
            self.emisor_cuit
            and self.cbte_tipo
            and self.punto_venta
            and self.numero
            and self.fecha
            and self.total is not None
            and self.cae
            and self.cae_vto
            and not self.inconsistente
        )


@dataclass
class NoReconocido:
    """El archivo no es una factura de ARCA que se pueda leer de forma determinística."""

    motivo: str
    texto: str = ""
    archivo: str = ""


@dataclass
class FilaMisComprobantes:
    fecha: date
    cbte_tipo: int
    punto_venta: int
    numero: int
    cae: str
    receptor_doc_tipo: int
    receptor_doc_nro: int
    receptor_nombre: str
    total: Decimal


@dataclass
class LecturaMisComprobantes:
    filas: list[FilaMisComprobantes]
    cuit_emisor: str = ""
    advertencias: list[str] = field(default_factory=list)


# --- utilidades -----------------------------------------------------------------------------


def _norm(t: str) -> str:
    sin = "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", sin).strip().lower()


def _fecha(t: str) -> date | None:
    m = re.search(r"(\d{2})/(\d{2})/(\d{4})", t)
    if not m:
        return None
    try:
        return date(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        return None


def _numero(t: str) -> Decimal:
    """«1093843,20», «1.093.843,20» o «1,00» -> Decimal."""
    t = t.strip().replace("$", "").replace(" ", "")
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    try:
        return Decimal(t)
    except InvalidOperation:
        raise ValueError(f"No es un número: {t!r}") from None


@dataclass(frozen=True)
class Corrida:
    x0: float
    x1: float
    y: float  # desde el borde superior
    texto: str

    @property
    def n(self) -> str:
        return _norm(self.texto)


def _corridas(page: Any) -> list[Corrida]:
    alto = page.get_size()[1]
    tp = page.get_textpage()
    salida = []
    for i in range(tp.count_rects()):
        izq, abajo, der, arriba = tp.get_rect(i)
        texto = tp.get_text_bounded(izq, abajo, der, arriba).replace("\r", " ").replace("\n", " ")
        if texto.strip():
            salida.append(Corrida(izq, der, alto - arriba, texto.strip()))
    return sorted(salida, key=lambda c: (round(c.y), c.x0))


def _derecha(rs: list[Corrida], etiqueta: Corrida, margen: float = 3.5) -> list[Corrida]:
    """Las corridas en la misma línea, a la derecha de la etiqueta."""
    return sorted(
        (
            c
            for c in rs
            if c is not etiqueta and abs(c.y - etiqueta.y) <= margen and c.x0 >= etiqueta.x1 - 1.5
        ),
        key=lambda c: c.x0,
    )


def _valor_derecha(rs: list[Corrida], etiqueta: Corrida) -> str:
    vs = _derecha(rs, etiqueta)
    return vs[0].texto if vs else ""


def _etiquetas(rs: list[Corrida], prefijo: str) -> list[Corrida]:
    p = _norm(prefijo)
    return [c for c in rs if c.n.startswith(p)]


# --- PDF ------------------------------------------------------------------------------------


def _render_qr(page: Any) -> dict[str, Any] | None:
    """Decodifica el QR de la página (imagen en los PDF de ARCA, vector en los de Factumono)."""
    try:
        import zxingcpp
    except ImportError:  # pragma: no cover - dependencia del paquete
        return None
    imagen = page.render(scale=3).to_pil()
    for r in zxingcpp.read_barcodes(imagen):
        if QR_PREFIJO in r.text:
            try:
                crudo = r.text.split("p=", 1)[1]
                crudo += "=" * (-len(crudo) % 4)
                datos = json.loads(base64.urlsafe_b64decode(crudo))
            except (ValueError, IndexError):
                continue
            if isinstance(datos, dict):
                return datos
    return None


COLUMNAS_TITULOS = (
    "codigo",
    "producto",
    "cantidad",
    "u. medida",
    "precio unit",
    "% bonif",
    "imp. bonif",
    "subtotal",
)


def leer_pdf(ruta: Path) -> FacturaLeida | NoReconocido:
    import pypdfium2 as pdfium

    nombre = ruta.name
    try:
        pdf = pdfium.PdfDocument(str(ruta))
    except Exception:
        return NoReconocido("No es un PDF válido.", archivo=nombre)
    try:
        if len(pdf) == 0:
            return NoReconocido("El PDF no tiene páginas.", archivo=nombre)
        paginas = []
        for i in range(len(pdf)):
            rs = _corridas(pdf[i])
            if i > 0 and any(c.n in ("duplicado", "triplicado") for c in rs[:6]):
                break  # las copias siguientes repiten lo mismo
            paginas.append(rs)
        texto_crudo = "\n".join(c.texto for rs in paginas for c in rs)
        if not texto_crudo.strip():
            return NoReconocido(
                "El PDF no tiene texto (es una imagen o un escaneo): no puedo leerlo de forma "
                "determinística.",
                archivo=nombre,
            )
        primera = paginas[0]
        if not _es_de_arca(primera):
            return NoReconocido(
                "No tiene el diseño de un comprobante de ARCA (faltan «CAE N°», «COD.» o «Comp. "
                "Nro»).",
                texto=texto_crudo,
                archivo=nombre,
            )
        f = _leer_cabecera(primera)
        f.archivo = nombre
        meta = pdf.get_metadata_dict() or {}
        creador = f"{meta.get('Creator', '')} {meta.get('Producer', '')}".lower()
        f.origen = "factumono" if "factumono" in creador else "arca"
        f.items = [it for rs in paginas for it in _leer_items(rs)]
        _leer_totales(primera, f)
        f.qr = _render_qr(pdf[0])
        _validar_con_qr(f)
        _controlar(f)
        return f
    finally:
        pdf.close()


def _es_de_arca(rs: list[Corrida]) -> bool:
    return (
        any(c.n.startswith("cae n") for c in rs)
        and any(re.match(r"cod\.\s*\d", c.n) for c in rs)
        and any(c.n.startswith("comp. nro") for c in rs)
    )


def _leer_cabecera(rs: list[Corrida]) -> FacturaLeida:
    f = FacturaLeida(origen="arca")
    # tipo de comprobante
    for c in rs:
        m = re.match(r"cod\.\s*(\d{1,3})", c.n)
        if m:
            f.cbte_tipo = int(m[1])
            break
    f.letra = next(
        (
            c.texto.strip()
            for c in rs[:12]
            if len(c.texto.strip()) == 1 and c.texto.strip() in "ABCEM"
        ),
        "",
    )
    for c in _etiquetas(rs, "punto de venta"):
        v = _valor_derecha(rs, c)
        if v.strip().isdigit():
            f.punto_venta = int(v)
    for c in _etiquetas(rs, "comp. nro"):
        v = _valor_derecha(rs, c)
        if v.strip().isdigit():
            f.numero = int(v)
    for c in _etiquetas(rs, "fecha de emision"):
        f.fecha = _fecha(_valor_derecha(rs, c))
    # emisor (columna izquierda arriba / derecha)
    for c in _etiquetas(rs, "razon social:"):
        if c.x0 < 100 and c.y < 130:
            f.emisor_nombre = _valor_derecha(rs, c)
    for c in _etiquetas(rs, "cuit:"):
        if c.x0 > 300 and c.y < 160:
            f.emisor_cuit = re.sub(r"\D", "", _valor_derecha(rs, c))
    for c in _etiquetas(rs, "ingresos brutos"):
        f.emisor_iibb = _valor_derecha(rs, c)
    for c in _etiquetas(rs, "fecha de inicio de actividades"):
        f.emisor_inicio_actividades = _fecha(_valor_derecha(rs, c))
    periodo = next(iter(_etiquetas(rs, "periodo facturado")), None)
    # el receptor arranca en «Apellido y Nombre / Razón Social:» (el período puede faltar)
    nombre_r = next(iter(_etiquetas(rs, "apellido y nombre")), None)
    y_periodo = (nombre_r.y - 6) if nombre_r else (periodo.y if periodo else 1e9)
    condiciones = _etiquetas(rs, "condicion frente al iva")
    emisor_cond = [c for c in condiciones if c.y <= y_periodo]
    receptor_cond = [c for c in condiciones if c.y > y_periodo]
    if emisor_cond:
        f.emisor_condicion_iva = _valor_derecha(rs, emisor_cond[0])
    # domicilio del emisor: las líneas debajo de «Domicilio Comercial:» hasta la etiqueta siguiente
    dom = next(iter(_etiquetas(rs, "domicilio comercial")), None)
    if dom:
        tope = emisor_cond[0].y if emisor_cond else dom.y + 40
        partes = [
            c.texto
            for c in rs
            if c.x0 >= dom.x1 - 1.5 and c.x1 < 335 and dom.y - 2 <= c.y < tope - 3
        ]
        f.emisor_domicilio = " ".join(partes)
    # período, vencimiento, condición de venta
    if periodo:
        f.periodo_desde = _fecha(_valor_derecha(rs, periodo))
        hasta = next(iter(_etiquetas(rs, "hasta:")), None)
        if hasta:
            f.periodo_hasta = _fecha(_valor_derecha(rs, hasta))
        vto = next(iter(_etiquetas(rs, "fecha de vto. para el pago")), None)
        if vto:
            f.vto_pago = _fecha(_valor_derecha(rs, vto))
    f.concepto = Concepto.SERVICIOS if f.periodo_desde else Concepto.PRODUCTOS
    cv = next(iter(_etiquetas(rs, "condicion de venta")), None)
    if cv:
        f.condicion_venta = _valor_derecha(rs, cv) or "Contado"
    # receptor
    nombre = next(iter(_etiquetas(rs, "apellido y nombre")), None)
    if nombre:
        f.receptor_nombre = _valor_derecha(rs, nombre)
    for c in rs:
        if c.y > y_periodo and c.x0 < 100 and re.match(r"(cuit|cuil|dni|doc)[^:]*:", c.n):
            valor = re.sub(r"\D", "", _valor_derecha(rs, c))
            if valor:
                f.receptor_doc_nro = int(valor)
                f.receptor_doc_tipo = (
                    int(DocTipo.DNI) if c.n.startswith("dni") else int(DocTipo.CUIT)
                )
            break
    if receptor_cond:
        f.receptor_condicion_iva = _valor_derecha(rs, receptor_cond[0])
    dom_r = next((c for c in _etiquetas(rs, "domicilio:") if c.x0 > 300 and c.y > y_periodo), None)
    if dom_r:
        cabecera = next((c for c in rs if c.n.startswith("codigo") and c.y > dom_r.y), None)
        tope = (cabecera.y - 5) if cabecera else dom_r.y + 40
        partes = [c.texto for c in rs if c.x0 >= dom_r.x1 - 1.5 and dom_r.y - 2 <= c.y < tope]
        f.receptor_domicilio = " ".join(partes)
    return f


def _leer_items(rs: list[Corrida]) -> list[ItemLeido]:
    cabecera = [c for c in rs if c.n.startswith("codigo") and c.y > 200]
    if not cabecera:
        return []
    h = cabecera[0]
    titulos = sorted((c for c in rs if abs(c.y - h.y) <= 3), key=lambda c: c.x0)
    arranques = [t.x0 for t in titulos]
    if len(arranques) != 8:
        return []  # otro diseño (por ejemplo A/B con IVA): no se lee por acá

    def columna(x: float) -> int:
        idx = 0
        for i, ini in enumerate(arranques):
            if x >= ini - 8:
                idx = i
        return idx

    # la tabla termina donde empieza el pie (totales, CAE, o «Pág.» si sigue en otra hoja)
    pie = re.compile(
        r"(subtotal:|importe otros|importe total|cae n|pag\.|comprobante autorizado"
        r"|esta (admin|agencia))"
    )
    fin = min((c.y for c in rs if c.y > h.y + 6 and pie.match(c.n)), default=1e9)
    zona = [c for c in rs if h.y + 6 < c.y < fin - 2]
    codigos = sorted((c for c in zona if columna(c.x0) == 0), key=lambda c: c.y)
    items: list[ItemLeido] = []
    for i, cod in enumerate(codigos):
        y0 = cod.y - 4
        y1 = codigos[i + 1].y - 4 if i + 1 < len(codigos) else fin - 2
        celdas: dict[int, list[Corrida]] = {}
        for c in zona:
            if y0 <= c.y < y1:
                celdas.setdefault(columna(c.x0), []).append(c)

        def texto(col: int, celdas: dict[int, list[Corrida]] = celdas) -> str:
            return " ".join(
                c.texto for c in sorted(celdas.get(col, []), key=lambda c: (round(c.y), c.x0))
            )

        try:
            items.append(
                ItemLeido(
                    codigo=cod.texto.strip(),
                    descripcion=texto(1),
                    cantidad=_numero(texto(2)),
                    unidad=texto(3),
                    precio_unitario=_numero(texto(4)),
                    bonificacion_pct=_numero(texto(5) or "0"),
                    bonificacion_importe=_numero(texto(6) or "0"),
                    subtotal=_numero(texto(7)),
                )
            )
        except ValueError:
            continue
    return items


def _leer_totales(rs: list[Corrida], f: FacturaLeida) -> None:
    for c in rs:
        if c.n.startswith("importe total"):
            with contextlib.suppress(ValueError):
                f.total = _numero(_valor_derecha(rs, c))
        elif c.n.startswith("cae n"):
            f.cae = re.sub(r"\D", "", _valor_derecha(rs, c))
        elif c.n.startswith("fecha de vto. de cae"):
            f.cae_vto = _fecha(_valor_derecha(rs, c))


def _validar_con_qr(f: FacturaLeida) -> None:
    """Compara los datos duros del texto con los del QR (RG 4892)."""
    if f.qr is None:
        f.advertencias.append("No pude leer el QR: los datos no se pudieron validar contra ARCA.")
        return
    q = f.qr
    fallas = []

    def comparar(nombre: str, texto: object, qr: object) -> None:
        if texto in (None, "", 0) or qr in (None, ""):
            return
        if str(texto) != str(qr):
            fallas.append(f"{nombre}: el texto dice {texto} y el QR dice {qr}")

    comparar("CUIT del emisor", f.emisor_cuit, q.get("cuit"))
    comparar("punto de venta", f.punto_venta, q.get("ptoVta"))
    comparar("número", f.numero, q.get("nroCmp"))
    comparar("tipo de comprobante", f.cbte_tipo, q.get("tipoCmp"))
    comparar("CAE", f.cae, q.get("codAut"))
    comparar("fecha", f.fecha.isoformat() if f.fecha else "", q.get("fecha"))
    if f.total is not None and q.get("importe") is not None:
        try:
            if Decimal(str(q["importe"])) != f.total:
                fallas.append(f"importe: el texto dice {f.total} y el QR dice {q['importe']}")
        except InvalidOperation:
            fallas.append("el importe del QR no es un número")
    if q.get("tipoDocRec") is not None:
        if f.receptor_doc_nro and (f.receptor_doc_tipo, f.receptor_doc_nro) != (
            q["tipoDocRec"],
            q.get("nroDocRec"),
        ):
            fallas.append(
                f"documento del receptor: el texto dice {f.receptor_doc_nro} y el QR dice "
                f"{q.get('nroDocRec')}"
            )
        elif not f.receptor_doc_nro:  # el texto no lo traía: se toma del QR
            f.receptor_doc_tipo = int(q["tipoDocRec"])
            f.receptor_doc_nro = int(q.get("nroDocRec") or 0)
    if fallas:
        f.inconsistente = True
        f.advertencias.append("El QR no coincide con el texto: " + "; ".join(fallas) + ".")
    # lo que el texto no trajo, el QR lo completa
    if not f.emisor_cuit and q.get("cuit"):
        f.emisor_cuit = str(q["cuit"])
    if not f.punto_venta and q.get("ptoVta"):
        f.punto_venta = int(q["ptoVta"])
    if not f.numero and q.get("nroCmp"):
        f.numero = int(q["nroCmp"])
    if not f.cbte_tipo and q.get("tipoCmp"):
        f.cbte_tipo = int(q["tipoCmp"])
    if f.total is None and q.get("importe") is not None:
        f.total = Decimal(str(q["importe"]))
    if not f.cae and q.get("codAut"):
        f.cae = str(q["codAut"])


def _controlar(f: FacturaLeida) -> None:
    """Controles de coherencia del texto (independientes del QR)."""
    if f.emisor_cuit and not cuit_mod.es_valido(f.emisor_cuit):
        f.advertencias.append(
            f"El CUIT del emisor ({f.emisor_cuit}) no tiene un dígito verificador válido."
        )
    if f.items:
        suma = sum((i.subtotal for i in f.items), Decimal(0))
        if f.total is not None and suma != f.total:
            f.advertencias.append(
                f"Los ítems suman {suma} pero el total es {f.total}: puede que falten ítems."
            )
    else:
        f.advertencias.append("No pude leer los ítems de la tabla.")
    for campo, valor in (
        ("emisor", f.emisor_nombre), ("fecha de emisión", f.fecha), ("receptor", f.receptor_nombre)
    ):  # fmt: skip
        if not valor:
            f.advertencias.append(f"No encontré el dato «{campo}».")


# --- «Mis Comprobantes» (Excel / CSV) ---------------------------------------------------------

_ENCABEZADOS = {
    "fecha": ("fecha", "fecha de emision", "fecha emision"),
    "tipo": ("tipo", "tipo de comprobante"),
    "pv": ("punto de venta", "pto. vta.", "pto vta"),
    "desde": ("numero desde", "nro. desde", "numero"),
    "hasta": ("numero hasta", "nro. hasta"),
    "cae": ("cod. autorizacion", "cod autorizacion", "cae", "codigo de autorizacion"),
    "doc_tipo": ("tipo doc. receptor", "tipo doc receptor", "tipo documento receptor"),
    "doc_nro": ("nro. doc. receptor", "nro doc receptor", "numero documento receptor"),
    "nombre": ("denominacion receptor", "denominacion", "razon social receptor"),
    "moneda": ("moneda",),
    "total": ("imp. total", "importe total", "total"),
}  # fmt: skip


def _clave_encabezado(t: str) -> str | None:
    n = _norm(t).rstrip(".")
    for clave, opciones in _ENCABEZADOS.items():
        if n in {o.rstrip(".") for o in opciones}:
            return clave
    return None


def _filas_crudas(ruta: Path) -> list[list[Any]]:
    if ruta.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        libro = load_workbook(ruta, read_only=True, data_only=True)
        hoja = libro.active
        assert hoja is not None
        filas = [list(f) for f in hoja.iter_rows(values_only=True)]
        libro.close()
        return filas
    crudo = ruta.read_bytes()
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = crudo.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    muestra = "\n".join(texto.splitlines()[:10])
    delimitador = max((";", ",", "\t"), key=muestra.count)
    return [list(f) for f in csv.reader(io.StringIO(texto), delimiter=delimitador)]


def _fecha_celda(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = str(v or "").strip()
    d = _fecha(t)
    if d:
        return d
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", t)
    return date(int(m[1]), int(m[2]), int(m[3])) if m else None


def _entero(v: Any) -> int | None:
    m = re.match(r"\s*(\d+)", str(v if v is not None else ""))
    return int(m[1]) if m else None


def _importe_celda(v: Any) -> Decimal:
    if isinstance(v, int | float | Decimal):
        return Decimal(str(v)).quantize(Decimal("0.01"))
    return _numero(str(v))


def leer_mis_comprobantes(ruta: Path) -> LecturaMisComprobantes:
    """Lee el Excel/CSV de «Mis Comprobantes → Emitidos». Tolera filas de título y distintos
    separadores; ignora lo que no sea en pesos y avisa de lo que no puede leer."""
    filas = _filas_crudas(ruta)
    cuit_titulo = ""
    for i, fila in enumerate(filas):  # noqa: B007 - `i` se usa después del for/else
        columnas = {_clave_encabezado(str(c)): j for j, c in enumerate(fila) if c is not None}
        columnas.pop(None, None)
        if {"fecha", "tipo", "pv", "total"} <= columnas.keys():
            break
        for celda in fila:
            m = re.search(r"\b(\d{2})-?(\d{8})-?(\d)\b", str(celda or ""))
            if m and not cuit_titulo:
                cuit_titulo = m[1] + m[2] + m[3]
    else:
        raise ValueError(
            "No reconozco el archivo: busco las columnas Fecha, Tipo, Punto de Venta e Imp. Total "
            "de «Mis Comprobantes → Emitidos»."
        )
    res = LecturaMisComprobantes([], cuit_titulo)
    for n, fila in enumerate(filas[i + 1 :], start=i + 2):
        if not any(str(c or "").strip() for c in fila):
            continue

        def celda(k: str, fila: list[Any] = fila) -> Any:
            j = columnas.get(k)
            return fila[j] if j is not None and j < len(fila) else None

        fecha = _fecha_celda(celda("fecha"))
        tipo, pv = _entero(celda("tipo")), _entero(celda("pv"))
        desde = _entero(celda("desde"))
        if fecha is None or tipo is None or pv is None or desde is None:
            res.advertencias.append(
                f"Fila {n}: no pude leerla (fecha, tipo, punto de venta o número)."
            )
            continue
        hasta = _entero(celda("hasta")) or desde
        if hasta != desde:
            res.advertencias.append(
                f"Fila {n}: es un rango de números ({desde}-{hasta}); la salteo."
            )
            continue
        moneda = str(celda("moneda") or "PES").strip().upper()
        if moneda not in ("PES", "$", "ARS", "PESOS"):
            res.advertencias.append(f"Fila {n}: moneda {moneda} (no es pesos); la salteo.")
            continue
        try:
            total = _importe_celda(celda("total"))
        except (ValueError, InvalidOperation):
            res.advertencias.append(f"Fila {n}: el importe total no es un número.")
            continue
        res.filas.append(
            FilaMisComprobantes(
                fecha=fecha,
                cbte_tipo=tipo,
                punto_venta=pv,
                numero=desde,
                cae=re.sub(r"\D", "", str(celda("cae") or "")),
                receptor_doc_tipo=_entero(celda("doc_tipo")) or int(DocTipo.CONSUMIDOR_FINAL),
                receptor_doc_nro=_entero(re.sub(r"\D", "", str(celda("doc_nro") or "")) or "0")
                or 0,
                receptor_nombre=str(celda("nombre") or "").strip() or "Consumidor Final",
                total=total,
            )
        )
    return res
